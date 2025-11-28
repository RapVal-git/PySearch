"""
Script de testare pentru verificarea suportului DOCX
"""
from docx import Document
import os

# Creează un fișier DOCX de test
def creeaza_docx_test():
    """Creează un fișier DOCX de test cu conținut sample"""
    doc = Document()
    
    # Adaugă titlu
    doc.add_heading('Document de Test pentru RAG', 0)
    
    # Adaugă paragrafe
    doc.add_paragraph('Acesta este un document de test pentru sistemul RAG.')
    doc.add_paragraph('Sistemul poate procesa atât fișiere PDF cât și DOCX.')
    
    doc.add_heading('Secțiunea 1: Informații Tehnice', level=1)
    doc.add_paragraph(
        'Python-docx este o bibliotecă puternică pentru procesarea documentelor Word. '
        'Permite extragerea textului, tabelelor și a altor elemente din fișiere DOCX.'
    )
    
    doc.add_heading('Secțiunea 2: Căutare Semantică', level=1)
    doc.add_paragraph(
        'Căutarea semantică folosește embeddings pentru a găsi documente relevante. '
        'Modelul paraphrase-multilingual-mpnet-base-v2 suportă multiple limbi, '
        'inclusiv română și engleză.'
    )
    
    # Adaugă un tabel
    doc.add_heading('Secțiunea 3: Tabel de Date', level=1)
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Light Grid Accent 1'
    
    # Header
    header_cells = table.rows[0].cells
    header_cells[0].text = 'Tip Document'
    header_cells[1].text = 'Extensie'
    header_cells[2].text = 'Suportat'
    
    # Date
    row1 = table.rows[1].cells
    row1[0].text = 'PDF'
    row1[1].text = '.pdf'
    row1[2].text = 'Da'
    
    row2 = table.rows[2].cells
    row2[0].text = 'Word'
    row2[1].text = '.docx'
    row2[2].text = 'Da'
    
    # Mai multe paragrafe pentru a testa paginarea
    doc.add_heading('Secțiunea 4: Conținut Extins', level=1)
    for i in range(15):
        doc.add_paragraph(
            f'Paragraf {i+1}: Acesta este un paragraf de test pentru a verifica '
            f'funcționalitatea de împărțire în "pagini" a documentelor DOCX. '
            f'Sistemul simulează pagini la fiecare 10 paragrafe.'
        )
    
    # Salvează documentul
    cale_test = 'test_document.docx'
    doc.save(cale_test)
    print(f"✅ Creat fișier de test: {cale_test}")
    return cale_test


def testeaza_extractie_docx():
    """Testează extracția de text din DOCX"""
    from docx_extractie import extrage_text_docx
    
    print("\n" + "="*60)
    print("TEST 1: Extracție Text din DOCX")
    print("="*60)
    
    cale_docx = creeaza_docx_test()
    
    if os.path.exists(cale_docx):
        cale_txt = extrage_text_docx(cale_docx)
        
        if cale_txt and os.path.exists(cale_txt):
            print(f"✅ Text extras cu succes în: {cale_txt}")
            
            # Afișează primele linii
            with open(cale_txt, 'r', encoding='utf-8') as f:
                lines = f.readlines()[:20]
                print("\n📄 Primele 20 linii din fișierul TXT:")
                print("".join(lines))
            
            return True
        else:
            print("❌ Eroare la extracția textului")
            return False
    else:
        print("❌ Fișierul DOCX nu a fost creat")
        return False


def testeaza_procesare_chunkuri():
    """Testează procesarea în chunk-uri"""
    from procesare_text import iterare_chunkuri_document
    
    print("\n" + "="*60)
    print("TEST 2: Procesare în Chunk-uri")
    print("="*60)
    
    cale_docx = 'test_document.docx'
    
    if not os.path.exists(cale_docx):
        print("⚠️  Creez fișier de test...")
        creeaza_docx_test()
    
    chunk_count = 0
    for item in iterare_chunkuri_document(cale_docx):
        chunk_count += 1
        if chunk_count <= 3:  # Afișează primele 3 chunk-uri
            print(f"\n--- Chunk {chunk_count} ---")
            print(f"Pagină: {item['metadata']['pagina']}")
            print(f"Tip: {item['metadata']['tip_document']}")
            print(f"Text: {item['text_chunk'][:200]}...")
    
    print(f"\n✅ Total chunk-uri procesate: {chunk_count}")
    return chunk_count > 0


def testeaza_indexare():
    """Testează indexarea unui document DOCX"""
    print("\n" + "="*60)
    print("TEST 3: Indexare Document DOCX")
    print("="*60)
    print("⚠️  Pentru acest test, rulează manual:")
    print("    python indexare_folder_qdrant.py")
    print("    și specifică un folder care conține test_document.docx")


if __name__ == "__main__":
    print("🧪 TESTARE SUPORT DOCX")
    print("="*60)
    
    # Rulează testele
    test1 = testeaza_extractie_docx()
    test2 = testeaza_procesare_chunkuri()
    testeaza_indexare()
    
    # Rezumat
    print("\n" + "="*60)
    print("📊 REZUMAT TESTE")
    print("="*60)
    print(f"Test 1 - Extracție DOCX: {'✅ PASS' if test1 else '❌ FAIL'}")
    print(f"Test 2 - Procesare Chunk-uri: {'✅ PASS' if test2 else '❌ FAIL'}")
    print("Test 3 - Indexare: ⚠️  Testare manuală necesară")
    
    if test1 and test2:
        print("\n🎉 Toate testele automate au trecut cu succes!")
        print("\n📝 Următorii pași:")
        print("1. Repornește API-ul pentru a încărca modificările")
        print("2. Rulează indexare_folder_qdrant.py pe un folder cu DOCX")
        print("3. Testează căutarea în documente DOCX")
    else:
        print("\n❌ Unele teste au eșuat. Verifică erorile de mai sus.")
